from pathlib import Path
from datetime import datetime

from docx import Document
from docx.shared import Pt


class ResumeGenerator:

    def __init__(self):

        self.output_folder = Path("output")
        self.output_folder.mkdir(exist_ok=True)

    def generate(self, tailored_resume, profile, job):

        document = Document()

        # ============================
        # Header
        # ============================

        title = document.add_heading(
            profile.get("name", ""),
            level=0
        )

        title.runs[0].font.size = Pt(22)

        document.add_paragraph(

            f'{profile.get("email","")} | '
            f'{profile.get("phone","")}'

        )

        document.add_paragraph(

            profile.get("location","")

        )

        document.add_paragraph(

            profile.get("linkedin","")

        )

        # ============================
        # Professional Summary
        # ============================

        document.add_heading(
            "Professional Summary",
            level=1
        )

        document.add_paragraph(

            tailored_resume.get(
                "summary",
                ""
            )

        )

        # ============================
        # Skills
        # ============================

        document.add_heading(
            "Skills",
            level=1
        )

        for skill in tailored_resume.get(
            "skills",
            []
        ):

            document.add_paragraph(
                skill,
                style="List Bullet"
            )

        # ============================
        # Experience
        # ============================

        document.add_heading(
            "Experience",
            level=1
        )

        for exp in tailored_resume.get(
            "experience",
            []
        ):

            document.add_heading(

                f'{exp.get("title","")} - '
                f'{exp.get("company","")}',

                level=2

            )

            bullets = (

                exp.get("bullets")

                or exp.get("description")

                or []

            )

            for bullet in bullets:

                document.add_paragraph(

                    bullet,

                    style="List Bullet"

                )

        # ============================
        # Education
        # ============================

        document.add_heading(
            "Education",
            level=1
        )

        for edu in profile.get(
            "education",
            []
        ):

            document.add_paragraph(

                f'{edu.get("degree","")}\n'

                f'{edu.get("school","")}'

            )

        # ============================
        # Certifications
        # ============================

        document.add_heading(
            "Certifications",
            level=1
        )

        for cert in profile.get(
            "certifications",
            []
        ):

            document.add_paragraph(

                cert,

                style="List Bullet"

            )

        # ============================
        # Target Job
        # ============================

        document.add_heading(
            "Target Position",
            level=1
        )

        document.add_paragraph(job.title)
        document.add_paragraph(job.company)

        # ============================
        # Footer
        # ============================

        document.add_heading(
            "Generated",
            level=1
        )

        document.add_paragraph(

            datetime.now().strftime(
                "%Y-%m-%d %H:%M"
            )

        )

        filename = self.output_folder / (

            f"{job.title.replace(' ','_')}_Resume.docx"

        )

        document.save(filename)

        return filename