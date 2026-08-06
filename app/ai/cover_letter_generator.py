from pathlib import Path
from datetime import datetime

from docx import Document
from docx.shared import Pt


class CoverLetterGenerator:

    def __init__(self):

        self.output = Path("output")

        self.output.mkdir(exist_ok=True)

    def generate(self, letter, profile, job):

        document = Document()

        title = document.add_heading(

            profile.get("name", ""),

            level=0

        )

        title.runs[0].font.size = Pt(22)

        document.add_paragraph(

            profile.get("email", "")

        )

        document.add_paragraph(

            profile.get("phone", "")

        )

        document.add_paragraph(

            profile.get("location", "")

        )

        document.add_heading(

            job.company,

            level=1

        )

        document.add_paragraph(

            f"Application for {job.title}"

        )

        document.add_paragraph(letter)

        document.add_paragraph("")

        document.add_paragraph(

            "Sincerely,"

        )

        document.add_paragraph(

            profile.get("name", "")

        )

        document.add_paragraph(

            datetime.now().strftime(
                "%Y-%m-%d"
            )

        )

        filename = self.output / (

            f"{job.title.replace(' ','_')}_CoverLetter.docx"

        )

        document.save(filename)

        return filename